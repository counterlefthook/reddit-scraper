"""Read Reddit URL lists from Excel/CSV and write scrape results to Excel."""

from __future__ import annotations

import csv
import re
from datetime import datetime, timezone
from pathlib import Path
from typing import Iterable

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .scraper import ThreadResult

_URL_RE = re.compile(
    r"https?://(?:[a-z]+\.)?(?:reddit\.com|redd\.it)/\S+|(?:^|\s)(?:www\.)?redd\.it/[a-z0-9]+",
    re.IGNORECASE,
)

# Excel cells reject control characters; strip anything openpyxl would choke on.
_ILLEGAL_XLSX_RE = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")

CELL_LIMIT = 32000  # Excel hard limit is 32767 chars per cell


def read_urls(path: str | Path) -> list[str]:
    """Extract Reddit URLs from an .xlsx, .csv, or .txt file.

    Scans every cell/line and pulls out anything that looks like a Reddit
    link, so the spreadsheet layout doesn't matter (headers, extra columns,
    notes alongside the links are all fine). Duplicates are dropped while
    preserving order.
    """
    path = Path(path)
    if not path.exists():
        raise FileNotFoundError(path)

    texts: list[str] = []
    suffix = path.suffix.lower()
    if suffix in (".xlsx", ".xlsm"):
        wb = load_workbook(path, read_only=True, data_only=True)
        try:
            for ws in wb.worksheets:
                for row in ws.iter_rows(values_only=True):
                    for cell in row:
                        if cell is not None:
                            texts.append(str(cell))
        finally:
            wb.close()
    elif suffix == ".csv":
        with open(path, newline="", encoding="utf-8-sig", errors="replace") as f:
            for row in csv.reader(f):
                texts.extend(row)
    else:  # treat anything else as plain text, one URL per line or free-form
        texts = Path(path).read_text(encoding="utf-8", errors="replace").splitlines()

    urls: list[str] = []
    seen: set[str] = set()
    for text in texts:
        for match in _URL_RE.findall(text):
            url = match.strip().rstrip('.,;)"\'')
            if url and url not in seen:
                seen.add(url)
                urls.append(url)
    return urls


def _clean(value):
    if isinstance(value, str):
        value = _ILLEGAL_XLSX_RE.sub("", value)
        if len(value) > CELL_LIMIT:
            value = value[:CELL_LIMIT] + "…[truncated]"
    return value


def _ts(epoch: float) -> str:
    if not epoch:
        return ""
    return datetime.fromtimestamp(epoch, tz=timezone.utc).strftime("%Y-%m-%d %H:%M:%S")


def _style_header(ws, widths: dict[int, int]) -> None:
    fill = PatternFill("solid", fgColor="DDEBF7")
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = fill
    ws.freeze_panes = "A2"
    for col, width in widths.items():
        ws.column_dimensions[get_column_letter(col)].width = width


POST_HEADERS = [
    "source_url", "post_id", "subreddit", "title", "author", "score",
    "upvote_ratio", "num_comments", "comments_scraped", "created_utc",
    "flair", "nsfw", "selftext", "permalink",
]

COMMENT_HEADERS = [
    "source_url", "post_id", "post_title", "comment_id", "parent_comment_id",
    "depth", "author", "is_op", "score", "created_utc", "body", "permalink",
]


def write_results(path: str | Path, results: Iterable[ThreadResult]) -> Path:
    """Write results to an .xlsx workbook with Posts, Comments, and Errors sheets."""
    results = list(results)
    wb = Workbook()

    posts_ws = wb.active
    posts_ws.title = "Posts"
    posts_ws.append(POST_HEADERS)
    comments_ws = wb.create_sheet("Comments")
    comments_ws.append(COMMENT_HEADERS)
    errors_ws = wb.create_sheet("Errors")
    errors_ws.append(["source_url", "error"])

    for res in results:
        if not res.ok:
            errors_ws.append([_clean(res.source_url), _clean(res.error or "unknown error")])
            continue
        p = res.post
        posts_ws.append([_clean(v) for v in [
            res.source_url, p.id, p.subreddit, p.title, p.author, p.score,
            p.upvote_ratio, p.num_comments, len(res.comments), _ts(p.created_utc),
            p.link_flair, "yes" if p.over_18 else "no", p.selftext, p.permalink,
        ]])
        for c in res.comments:
            comments_ws.append([_clean(v) for v in [
                res.source_url, c.post_id, p.title, c.id, c.parent_id, c.depth,
                c.author, "yes" if c.is_submitter else "no", c.score,
                _ts(c.created_utc), c.body, c.permalink,
            ]])

    _style_header(posts_ws, {1: 40, 3: 16, 4: 50, 5: 18, 13: 60, 14: 40})
    _style_header(comments_ws, {1: 40, 3: 40, 7: 18, 11: 80, 12: 40})
    _style_header(errors_ws, {1: 50, 2: 60})
    for row in comments_ws.iter_rows(min_row=2, min_col=11, max_col=11):
        row[0].alignment = Alignment(wrap_text=False, vertical="top")

    path = Path(path)
    wb.save(path)
    return path
