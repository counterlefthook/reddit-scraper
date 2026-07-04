from openpyxl import Workbook, load_workbook

from reddit_scraper.excel_io import read_urls, write_results
from reddit_scraper.scraper import Comment, Post, ThreadResult


def make_xlsx(path, rows, sheet2_rows=None):
    wb = Workbook()
    ws = wb.active
    for row in rows:
        ws.append(row)
    if sheet2_rows:
        ws2 = wb.create_sheet("More")
        for row in sheet2_rows:
            ws2.append(row)
    wb.save(path)


def test_read_urls_from_xlsx_any_cell_any_sheet(tmp_path):
    path = tmp_path / "urls.xlsx"
    make_xlsx(
        path,
        rows=[
            ["My research links", None],
            ["https://www.reddit.com/r/a/comments/111aaa/x/", "note about it"],
            ["see also https://redd.it/222bbb please", None],
            ["https://www.reddit.com/r/a/comments/111aaa/x/", "duplicate"],
            ["https://example.com/not-reddit", None],
        ],
        sheet2_rows=[["https://old.reddit.com/r/b/comments/333ccc/y/"]],
    )
    urls = read_urls(path)
    assert urls == [
        "https://www.reddit.com/r/a/comments/111aaa/x/",
        "https://redd.it/222bbb",
        "https://old.reddit.com/r/b/comments/333ccc/y/",
    ]


def test_read_urls_from_csv(tmp_path):
    path = tmp_path / "urls.csv"
    path.write_text(
        "url,notes\nhttps://www.reddit.com/r/a/comments/111aaa/x/,first\n"
    )
    assert read_urls(path) == ["https://www.reddit.com/r/a/comments/111aaa/x/"]


def sample_result():
    post = Post(
        id="1abcde", url="https://redd.it/1abcde", subreddit="AskReddit",
        title="A title", author="op", score=10, upvote_ratio=0.9,
        num_comments=2, created_utc=1735689600.0, selftext="body \x01 with control char",
        link_flair="", is_video=False, over_18=False,
        permalink="https://www.reddit.com/r/AskReddit/comments/1abcde/",
    )
    comments = [
        Comment(id="c1", post_id="1abcde", parent_id="", depth=0, author="u1",
                body="hello", score=5, created_utc=1735693200.0,
                is_submitter=False, stickied=False, permalink="https://reddit.com/c1"),
    ]
    return ThreadResult(source_url=post.url, post=post, comments=comments)


def test_write_results_roundtrip(tmp_path):
    out = tmp_path / "results.xlsx"
    failed = ThreadResult(source_url="https://redd.it/broken", error="thread not found (404)")
    write_results(out, [sample_result(), failed])

    wb = load_workbook(out)
    assert wb.sheetnames == ["Posts", "Comments", "Errors"]

    posts = list(wb["Posts"].values)
    assert posts[0][0] == "source_url"
    assert posts[1][3] == "A title"
    assert "\x01" not in posts[1][12]  # control chars stripped

    comments = list(wb["Comments"].values)
    assert comments[1][3] == "c1"
    assert comments[1][10] == "hello"

    errors = list(wb["Errors"].values)
    assert errors[1] == ("https://redd.it/broken", "thread not found (404)")
